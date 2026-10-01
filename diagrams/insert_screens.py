# -*- coding: utf-8 -*-
"""Вставляет скриншоты интерфейса в отчёт вместо таблиц-плейсхолдеров.
Берёт ОТЧЕТ_ПП_Галамарт_ФИНАЛ.docx (там уже 5 диаграмм) и добавляет 15 скринов.
"""
import struct
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm

SRC = Path(r"D:\TG_ZAGRUZKI\ОТЧЕТ_ПП_Галамарт_ФИНАЛ.docx")
OUT = Path(r"D:\TG_ZAGRUZKI\ОТЧЕТ_ПП_Галамарт_ФИНАЛ_со_скринами.docx")
IMG = Path(r"C:\Users\user\Desktop\практичке\pp11\diagrams\screens")

MAPPING = {
    "Рисунок 3.1": IMG / "fig_3_1_dashboard.png",
    "Рисунок 3.2": IMG / "fig_3_2_products.png",
    "Рисунок 3.3": IMG / "fig_3_3_orders.png",
    "Рисунок 3.4": IMG / "fig_3_4_mobile.png",
    "Рисунок 4.1": IMG / "fig_4_1_login.png",
    "Рисунок 4.2": IMG / "fig_4_2_dashboard_admin.png",
    "Рисунок 4.3": IMG / "fig_4_3_product_add.png",
    "Рисунок 4.4": IMG / "fig_4_4_order_form.png",
    "Рисунок 4.5": IMG / "fig_4_5_deficit_report.png",
    "Рисунок 4.6": IMG / "fig_4_6_pdf_export.png",
    "Рисунок 4.7": IMG / "fig_4_7_validation.png",
    "Рисунок А.1": IMG / "fig_A_1_login_steps.png",
    "Рисунок А.2": IMG / "fig_A_2_order_steps.png",
    "Рисунок Г.1": IMG / "fig_G_1_input_card.png",
    "Рисунок Г.2": IMG / "fig_G_2_report_data.png",
}


def png_size(path):
    with open(path, "rb") as f:
        head = f.read(24)
    w, h = struct.unpack(">II", head[16:24])
    return w, h


def cell_text(tbl):
    return " ".join(
        "".join(t.text or "" for t in c.iter(qn("w:t")))
        for r in tbl.findall(qn("w:tr"))
        for c in r.findall(qn("w:tc"))
    )


def para_text(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()


def main():
    doc = Document(SRC)
    body = doc.element.body
    children = list(body)
    inserted = 0
    for i, el in enumerate(children):
        if el.tag != qn("w:tbl"):
            continue
        if "МЕСТО ДЛЯ СКРИНШОТА" not in cell_text(el):
            continue
        cap = None
        for j in range(i + 1, min(i + 4, len(children))):
            if children[j].tag == qn("w:p"):
                tx = para_text(children[j])
                if tx:
                    cap = tx
                    break
        img = None
        for key, path in MAPPING.items():
            if cap and cap.startswith(key):
                img = path
                break
        if img is None:
            continue
        w, h = png_size(img)
        width = Cm(15.5) if w >= h else Cm(10.0)
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Cm(0.2)
        p.paragraph_format.space_after = Cm(0.2)
        p.add_run().add_picture(str(img), width=width)
        el.addprevious(p._p)
        body.remove(el)
        inserted += 1
        print("inserted", cap[:55], "->", img.name)

    doc.save(OUT)
    left = sum(1 for t in doc.tables if "МЕСТО ДЛЯ СКРИНШОТА" in cell_text(t._tbl))
    print("saved", OUT)
    print("inserted", inserted, "| placeholders left", left)


if __name__ == "__main__":
    main()
