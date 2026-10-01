# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.shared import Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

SRC = Path(r"D:\TG_ZAGRUZKI\ОТЧЕТ_ПП_Галамарт_ГОТОВЫЙ_без_лишних_таблиц.docx")
OUT = Path(r"D:\TG_ZAGRUZKI\ОТЧЕТ_ПП_Галамарт_ФИНАЛ.docx")
IMG = Path(r"C:\Users\user\Desktop\практичке\pp11\diagrams")

MAPPING = {
    "Рисунок 1.1": IMG / "idef0_asis.png",
    "Рисунок 2.1": IMG / "idef0_system.png",
    "Рисунок 2.2": IMG / "usecase.png",
    "Рисунок 2.3": IMG / "sequence.png",
    "Рисунок 2.4": IMG / "er_database.png",
}


def cell_text(tbl):
    return " ".join("".join(t.text or "" for t in c.iter(qn("w:t")))
                    for r in tbl.findall(qn("w:tr")) for c in r.findall(qn("w:tc")))


def para_text(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()


doc = Document(SRC)
body = doc.element.body
children = list(body)
inserted = 0
for i, el in enumerate(children):
    if el.tag != qn("w:tbl"):
        continue
    if "МЕСТО ДЛЯ СКРИНШОТА" not in cell_text(el):
        continue
    cap = None
    for j in range(i + 1, min(i + 4, len(children))):
        if children[j].tag == qn("w:p"):
            tx = para_text(children[j])
            if tx:
                cap = tx
                break
    img = None
    for key, path in MAPPING.items():
        if cap and cap.startswith(key):
            img = path
            break
    if img is None:
        continue
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Cm(0.2)
    p.paragraph_format.space_after = Cm(0.2)
    run = p.add_run()
    run.add_picture(str(img), width=Cm(15.5))
    el.addprevious(p._p)
    body.remove(el)
    inserted += 1

doc.save(OUT)
print("saved", OUT)
print("inserted", inserted, "tables_left", len(doc.tables))
